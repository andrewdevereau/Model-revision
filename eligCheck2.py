# -*- coding: utf-8 -*-
"""
Created on Mon Jul 11 14:30:15 2016

@author: Andrew Devereau
takes the eligibiltiy word docs from the new model catalogue and does a 
comparison producing a html table
"""

from docx import Document
from docx.enum.text import WD_BREAK
from difflib import HtmlDiff

def getData (file):
    #gets disease eligibility text from tabulated statements and makes a list
    document = Document(file)
    diseaseList = []
    tables = document.tables
    for table in tables:
        if table.cell(0,0).text == 'Level 3 Title':
            disease = {}
            title,x,y = table.cell(1,1).text.rpartition('(')
            disease['Title'] = title.strip()
            disease['Text'] = table.cell(2,1).text
            diseaseList.append(disease)
    return diseaseList

def compareDocs(newList, oldList):
    #uses html compare to generate a comparison of two eligibility docs
    k = HtmlDiff(wrapcolumn=80)   #gets an instance of html diff with wrapped columns
    with open('diff.html', 'w') as output:  #a file for the output

        for f in newList:           #go through each disease in new
            for g in oldList:       #search through each disease in ref
                if f['Title'] == g['Title']:
                    text1 = g['Text']
                    text2 = f['Text']
                    diffFile = k.make_file(text1.splitlines(keepends = True), text2.splitlines(keepends=True),fromdesc='old Model', todesc='new Model', charset='ISO-8859-1')
                    output.write(diffFile)  #write the Diff to the file
                    break
            else:
                print(f['Title'], " has no match")
    return
    
#main
newList = getData('Rare Disease Eligibility Criteria - v1.6.0.docx')
oldList = getData('Rare Disease Eligibility Criteria - v1.5.1 - FINAL.docx')
compareDocs(newList, oldList)

    
   

    
    
  
    
   
  
